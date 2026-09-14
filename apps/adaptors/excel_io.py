"""Extract header-keyed rows from Excel (.xlsx) workbooks."""

from __future__ import annotations

import io
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, BinaryIO, Iterable, Union

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

SourceInput = Union[str, Path, bytes, BinaryIO]


def _cell_as_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S") if (value.hour or value.minute or value.second) else value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        return format(Decimal(str(value)).normalize(), "f").rstrip("0").rstrip(".") or "0"
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def _header_label(raw: Any, index: int) -> str:
    text = _cell_as_str(raw)
    return text if text else f"column_{index + 1}"


def _row_is_empty(values: tuple[Any, ...]) -> bool:
    return all(v is None or str(v).strip() == "" for v in values)


def _active_worksheet(workbook: Workbook) -> Worksheet:
    sheet = workbook.active
    if sheet is not None:
        return sheet
    sheets = workbook.worksheets
    if not sheets:
        raise ValueError("Excel workbook has no worksheets.")
    return sheets[0]


def iter_excel_dict_rows(file_stream: BinaryIO) -> Iterable[dict[str, Any]]:
    workbook = load_workbook(file_stream, read_only=True, data_only=True)
    worksheet = _active_worksheet(workbook)
    rows = worksheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        workbook.close()
        return

    headers = [_header_label(cell, i) for i, cell in enumerate(header_row)]
    for row in rows:
        if row is None or _row_is_empty(row):
            continue
        padded = list(row) + [None] * max(0, len(headers) - len(row))
        yield {headers[i]: _cell_as_str(padded[i]) for i in range(len(headers))}
    workbook.close()


def extract_excel_rows(source_input: SourceInput) -> Iterable[dict[str, Any]]:
    """Yield header-keyed rows from an .xlsx path or byte stream."""
    if isinstance(source_input, (str, Path)):
        with open(source_input, mode="rb") as handle:
            for row in iter_excel_dict_rows(handle):
                yield row
        return
    if isinstance(source_input, bytes):
        for row in iter_excel_dict_rows(io.BytesIO(source_input)):
            yield row
        return
    if hasattr(source_input, "read"):
        content = source_input.read()
        if not isinstance(content, bytes):
            raise ValueError("Excel source stream must yield bytes.")
        for row in iter_excel_dict_rows(io.BytesIO(content)):
            yield row
        return
    raise ValueError(f"Unsupported excel source_input type: {type(source_input)}")
